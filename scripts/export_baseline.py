#!/usr/bin/env python3.11
# -*- coding: utf-8 -*-
"""
export_baseline.py — 从报价表 xlsx 导出「文本报价基线」到 assets/data/nanjing_baseline.json

作用：把南京仓报价基线（xlsx「蔬菜」表）转成纯文本 JSON，
      使仓库可纯文本推送（GitHub MCP / git 无需处理二进制），
      并让每日工作流不再依赖 openpyxl 解析 xlsx，更稳健。

当报价表 xlsx 更新时，重跑一次本脚本刷新基线即可：
    python3.11 scripts/export_baseline.py
"""
import json
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(ROOT, "马厂老火锅10月份上半月蔬菜报价.xlsx")
OUT = os.path.join(ROOT, "assets", "data", "nanjing_baseline.json")


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["蔬菜"]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, 9)]
        if not row or not row[1]:
            continue
        _id, name, sys_name, unit, quote, standard, cur_price, source = row
        rows.append({
            "id": int(_id) if isinstance(_id, (int, float)) else None,
            "name": str(name).strip(),
            "sysName": (str(sys_name).strip() if sys_name else ""),
            "unit": (str(unit).strip() if unit else "斤"),
            "quote": quote if isinstance(quote, (int, float)) else None,
            "standard": (str(standard).strip() if standard else ""),
            "curPrice": cur_price if isinstance(cur_price, (int, float)) else None,
            "source": (str(source).strip() if source else ""),
        })
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump({"generatedFrom": os.path.basename(XLSX), "rowCount": len(rows), "rows": rows},
                  f, ensure_ascii=False, indent=2)
    print(f"✅ 导出 {len(rows)} 行 -> {OUT}")


if __name__ == "__main__":
    main()
